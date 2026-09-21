from __future__ import annotations

import csv
import os
from pathlib import Path

from src.env_bootstrap import PROJECT_ROOT
from src.schemas import ContactLead, LeadResearchReport

DEFAULT_LEADS_CSV = PROJECT_ROOT / "leads.csv"
DEFAULT_LEADS_XLSX = PROJECT_ROOT / "leads.xlsx"

CSV_COLUMNS = [
    "company",
    "area",
    "vertical",
    "generated_utc",
    "phones",
    "emails",
    "linkedin",
    "websites",
    "other_contacts",
    "company_age_signal",
    "executive_summary",
    "hypotheses",
    "discovery_questions",
    "talking_points",
    "sources",
    "tool_notes",
    "status",
    "your_notes",
]


def _join(values: list[str], *, sep: str = " | ") -> str:
    cleaned = [v.strip() for v in values if v and v.strip()]
    return sep.join(cleaned)


def _contacts_by_kind(contacts: list[ContactLead]) -> dict[str, list[str]]:
    buckets: dict[str, list[str]] = {
        "phone": [],
        "email": [],
        "linkedin": [],
        "website": [],
        "other": [],
    }
    for c in contacts:
        kind = (c.kind or "other").strip().lower()
        val = (c.value or "").strip()
        if not val:
            continue
        key = kind if kind in buckets else "other"
        if val not in buckets[key]:
            buckets[key].append(val)
    return buckets


def report_to_csv_row(report: LeadResearchReport) -> dict[str, str]:
    brief = report.brief
    contacts = _contacts_by_kind(report.contact_leads)

    sources = _join(
        [f"{s.title} — {s.url}" for s in report.sources[:12]],
        sep=" ;; ",
    )
    if len(report.sources) > 12:
        sources = f"{sources} ;; … (+{len(report.sources) - 12} more)"

    exec_summary = (brief.executive_summary if brief else "").strip()
    if not exec_summary and report.sources:
        top = report.sources[0]
        exec_summary = f"(No AI brief) {top.title[:120]} — see sources column."
    elif not exec_summary:
        exec_summary = "(No AI brief — set LLM credentials in .env; avoid --no-llm for full rows.)"

    age_signal = (brief.company_age_signal if brief else "").strip()
    if not age_signal and report.playbook:
        age_signal = "(Checklist only — run without --no-llm for AI startup-age signal.)"

    phones = _join(contacts["phone"])
    emails = _join(contacts["email"])
    websites = _join(contacts["website"])
    other = _join(contacts["other"])
    if not phones and not emails:
        # Ensure spreadsheet has a clickable site when outreach details are missing.
        site_links = websites or other
        if not site_links and report.sources:
            from src.contacts import extract_company_websites

            site_links = _join([u for u, _ in extract_company_websites(report.sources, report.company)])
        if site_links and not other:
            other = site_links

    notes_parts: list[str] = []
    if report.notes:
        notes_parts.append(report.notes.strip())
    if not brief and report.sources:
        notes_parts.append(f"{len(report.sources)} web source(s) captured.")
    if not phones and not emails:
        notes_parts.append("No phone/email in snippets — website link in other_contacts.")

    return {
        "company": report.company,
        "area": report.locality or "",
        "vertical": report.industry or "",
        "generated_utc": report.generated_at.strftime("%Y-%m-%d %H:%M UTC"),
        "phones": phones,
        "emails": emails,
        "linkedin": _join(contacts["linkedin"]),
        "websites": websites,
        "other_contacts": other,
        "company_age_signal": age_signal,
        "executive_summary": exec_summary,
        "hypotheses": _join(brief.hypotheses if brief else []),
        "discovery_questions": _join(brief.discovery_questions if brief else []),
        "talking_points": _join(brief.talking_points if brief else []),
        "sources": sources,
        "tool_notes": " | ".join(notes_parts),
        "status": "",
        "your_notes": "",
    }


def _canonical_header(cell_value: object) -> str | None:
    if cell_value is None:
        return None
    key = str(cell_value).strip().lower().replace(" ", "_")
    if key in CSV_COLUMNS:
        return key
    for col in CSV_COLUMNS:
        if key == col.upper():
            return col
    return None


def default_xlsx_path() -> Path:
    """Project leads.xlsx, or LEADS_XLSX_PATH from .env (e.g. a Teams/OneDrive workbook)."""
    env_path = (os.environ.get("LEADS_XLSX_PATH") or "").strip().strip('"').strip("'")
    if env_path:
        p = Path(env_path)
        return p if p.is_absolute() else PROJECT_ROOT / p
    return DEFAULT_LEADS_XLSX


def resolve_xlsx_path(path: Path | str | None) -> Path:
    if path is None:
        return default_xlsx_path()
    p = Path(path)
    if not p.is_absolute():
        p = PROJECT_ROOT / p
    return p


def resolve_csv_path(path: Path | str | None) -> Path:
    """Relative paths are under the project folder (not the shell cwd)."""
    if path is None:
        return DEFAULT_LEADS_CSV
    p = Path(path)
    if not p.is_absolute():
        p = PROJECT_ROOT / p
    return p


def read_companies_from_xlsx(path: Path | str | None = None) -> list[str]:
    out = resolve_xlsx_path(path)
    if not out.exists():
        return []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    wb = load_workbook(out, read_only=True, data_only=True)
    ws = wb.active
    header_row, header_map = _sheet_header_map(ws)
    company_col = header_map.get("company", 1)
    names: list[str] = []
    max_row = getattr(ws, "max_row", None) or header_row
    for row in range(header_row + 1, max_row + 1):
        val = ws.cell(row, company_col).value
        if val is None:
            continue
        company = str(val).strip()
        if company and company.lower() not in {"quicktest", "unknown", "company"}:
            names.append(company)
    wb.close()
    return names


def read_companies_from_spreadsheet(
    *,
    xlsx_path: Path | str | None = None,
    csv_path: Path | str | None = None,
) -> list[str]:
    """Prefer leads.xlsx (formatted workbook), fall back to leads.csv."""
    names = read_companies_from_xlsx(xlsx_path)
    if names:
        return names
    return read_companies_from_csv(csv_path)


def read_companies_from_csv(path: Path | str | None = None) -> list[str]:
    """Company names already saved in the spreadsheet (for --exclude-seen)."""
    out = resolve_csv_path(path)
    if not out.exists() or out.stat().st_size == 0:
        return []
    names: list[str] = []
    with out.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            company = (row.get("company") or "").strip()
            if company and company.lower() not in {"quicktest", "unknown"}:
                names.append(company)
    return names


def count_csv_data_rows(path: Path) -> int:
    if not path.exists() or path.stat().st_size == 0:
        return 0
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        rows = list(reader)
    return max(0, len(rows) - 1)  # minus header


def append_report_to_csv(
    report: LeadResearchReport,
    path: Path | str,
    *,
    append: bool = True,
) -> Path:
    """Write one report row to CSV. Creates file + header on first write."""
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    row = report_to_csv_row(report)

    file_exists = out.exists() and out.stat().st_size > 0
    write_header = not file_exists or not append
    mode = "a" if append and file_exists else "w"

    with out.open(mode, encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_COLUMNS, extrasaction="ignore")
        if write_header:
            writer.writeheader()
        writer.writerow(row)

    return out.resolve()


def _sheet_header_map(ws) -> tuple[int, dict[str, int]]:
    """Works with normal and read-only worksheets from openpyxl."""
    max_row = getattr(ws, "max_row", None) or 1
    max_column = getattr(ws, "max_column", None) or len(CSV_COLUMNS)
    for header_row in range(1, min(5, max_row + 1)):
        mapping: dict[str, int] = {}
        for col in range(1, max_column + 1):
            canonical = _canonical_header(ws.cell(header_row, col).value)
            if canonical and canonical not in mapping:
                mapping[canonical] = col
        if "company" in mapping:
            return header_row, mapping
    return 1, {name: idx + 1 for idx, name in enumerate(CSV_COLUMNS)}


def _next_data_row(ws, header_row: int, company_col: int) -> int:
    max_row = getattr(ws, "max_row", None) or header_row
    last = header_row
    for row in range(header_row + 1, max_row + 2):
        val = ws.cell(row, company_col).value
        if val is None:
            continue
        text = str(val).strip()
        if not text:
            continue
        if text.lower() in {"company", "quicktest"}:
            continue
        if _canonical_header(text) == "company":
            continue
        last = row
    return last + 1


def count_xlsx_data_rows(path: Path) -> int:
    if not path.exists():
        return 0
    try:
        from openpyxl import load_workbook
    except ImportError:
        return 0
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header_row, header_map = _sheet_header_map(ws)
    company_col = header_map.get("company", 1)
    count = 0
    max_row = getattr(ws, "max_row", None) or header_row
    for row in range(header_row + 1, max_row + 1):
        val = ws.cell(row, company_col).value
        if val and str(val).strip() and str(val).strip().lower() not in {"quicktest", "company"}:
            count += 1
    wb.close()
    return count


def append_report_to_xlsx(
    report: LeadResearchReport,
    path: Path | str,
    *,
    append: bool = True,
) -> Path:
    """Append one row to .xlsx without changing header row formatting."""
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl: .venv\\Scripts\\pip install openpyxl") from exc

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    row_data = report_to_csv_row(report)

    if out.exists() and append:
        wb = load_workbook(out)
        ws = wb.active
        header_row, header_map = _sheet_header_map(ws)
        company_col = header_map.get("company", 1)
        target_row = _next_data_row(ws, header_row, company_col)
        for field, col_idx in header_map.items():
            ws.cell(target_row, col_idx, row_data.get(field, ""))
        wb.save(out)
        wb.close()
        return out.resolve()

    if out.exists() and not append:
        wb = load_workbook(out)
        ws = wb.active
        header_row, header_map = _sheet_header_map(ws)
        if ws.max_row > header_row:
            ws.delete_rows(header_row + 1, ws.max_row - header_row)
        target_row = header_row + 1
        for field, col_idx in header_map.items():
            ws.cell(target_row, col_idx, row_data.get(field, ""))
        wb.save(out)
        wb.close()
        return out.resolve()

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    for idx, name in enumerate(CSV_COLUMNS, start=1):
        ws.cell(1, idx, name)
    for idx, name in enumerate(CSV_COLUMNS, start=1):
        ws.cell(2, idx, row_data.get(name, ""))
    wb.save(out)
    wb.close()
    return out.resolve()
